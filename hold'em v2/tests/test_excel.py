"""Excel export tests - append, duplicate prevention and rebuild."""

import datetime

import pytest
from openpyxl import load_workbook

from export import excel_export

ROW = {
    "id": 1,
    "recorded_at": datetime.datetime(2026, 9, 2, 12, 30, 0),
    "player_card_1": "8S", "player_card_2": "9C",
    "flop_card_1": "2D", "flop_card_2": "QH", "flop_card_3": "3D",
    "turn_card": "3H", "river_card": "6S",
    "dealer_card_1": "8D", "dealer_card_2": "QD",
    "player_hand": "Pair", "dealer_hand": "Two Pair",
    "winner": "Dealer", "dealer_qualified": True,
}


@pytest.fixture
def workbook_path(tmp_path):
    return str(tmp_path / "poker_hands.xlsx")


def read_rows(path):
    workbook = load_workbook(path)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    return rows


def read_columns(path, index=1):
    """One row as {column heading: value}, so tests do not depend on order."""
    rows = read_rows(path)
    return dict(zip(rows[0], rows[index]))


def test_creates_workbook_with_the_documented_headers(workbook_path):
    excel_export.ensure_workbook(workbook_path)
    assert list(read_rows(workbook_path)[0]) == excel_export.HEADERS


def test_appends_a_hand(workbook_path):
    assert excel_export.append_hand(ROW, workbook_path) is True
    rows = read_rows(workbook_path)
    assert len(rows) == 2
    assert rows[1][0] == 1
    assert rows[1][2:4] == ("8S", "9C")
    columns = read_columns(workbook_path)
    assert columns["Player Hand"] == "Pair"
    assert columns["Dealer Hand"] == "Two Pair"
    assert columns["Winner"] == "Dealer"
    assert columns["Dealer Qualified"] == "Yes"


def test_does_not_append_the_same_hand_twice(workbook_path):
    excel_export.append_hand(ROW, workbook_path)
    assert excel_export.append_hand(ROW, workbook_path) is False
    assert len(read_rows(workbook_path)) == 2


def test_appends_a_different_hand(workbook_path):
    excel_export.append_hand(ROW, workbook_path)
    second = dict(ROW, id=2, river_card="7S")
    assert excel_export.append_hand(second, workbook_path) is True
    assert len(read_rows(workbook_path)) == 3


def test_rebuild_replaces_the_file(workbook_path):
    excel_export.append_hand(ROW, workbook_path)
    count = excel_export.rebuild([ROW, dict(ROW, id=2)], workbook_path)
    rows = read_rows(workbook_path)
    assert count == 2
    assert len(rows) == 3
    assert [row[0] for row in rows[1:]] == [1, 2]


def test_rebuild_recreates_a_deleted_file(tmp_path):
    path = str(tmp_path / "gone.xlsx")
    excel_export.rebuild([ROW], path)
    assert read_rows(path)[1][0] == 1


def test_timestamp_is_written_readably(workbook_path):
    excel_export.append_hand(ROW, workbook_path)
    assert read_rows(workbook_path)[1][1] == "2026-09-02 12:30:00"


def test_dealer_qualification_is_written_as_yes_or_no(workbook_path):
    excel_export.append_hand(dict(ROW, dealer_qualified=False), workbook_path)
    assert read_columns(workbook_path)["Dealer Qualified"] == "No"


def test_a_hand_from_before_the_winner_column_leaves_it_blank(workbook_path):
    older = {key: value for key, value in ROW.items()
             if key not in ("winner", "dealer_qualified")}
    excel_export.append_hand(older, workbook_path)
    columns = read_columns(workbook_path)
    assert columns["Winner"] is None
    assert columns["Dealer Qualified"] is None


def test_a_hand_recorded_without_timings_leaves_them_blank(workbook_path):
    """Hands stored before the timer existed still export cleanly."""
    excel_export.append_hand(ROW, workbook_path)
    columns = read_columns(workbook_path)
    assert columns["Round Length (s)"] is None
    assert columns["Since Previous Round (s)"] is None


def test_locked_file_raises_a_clear_error(workbook_path, monkeypatch):
    def refuse(*_args, **_kwargs):
        raise PermissionError("file is open in Excel")

    excel_export.ensure_workbook(workbook_path)
    monkeypatch.setattr(excel_export.Workbook, "save", refuse)
    with pytest.raises(excel_export.ExcelError) as error:
        excel_export.rebuild([ROW], workbook_path)
    assert "close the file in Excel" in str(error.value).lower() or "cannot write" in str(error.value).lower()
