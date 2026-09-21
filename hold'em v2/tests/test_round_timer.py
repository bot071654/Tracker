"""Round timings: how long a round runs, and how long until the next one."""

import pytest

from tracker import (
    COMPLETE, FLOP, PLAYER_CARDS, RIVER, RoundTimer, TURN, WAITING,
)


class FakeClock:
    """A clock that only moves when the test says so."""

    def __init__(self):
        self.now = 1000.0

    def __call__(self):
        return self.now

    def advance(self, seconds):
        self.now += seconds
        return self.now


@pytest.fixture
def clock():
    return FakeClock()


@pytest.fixture
def timer(clock):
    return RoundTimer(clock=clock, wall_clock=lambda: "wall-%.1f" % clock.now)


def play_a_round(timer, clock, deal_to_flop=6.0, flop_to_showdown=9.0):
    """Walk one round through the states the tracker reports."""
    timer.observe(PLAYER_CARDS)
    clock.advance(deal_to_flop)
    timer.observe(FLOP)
    clock.advance(flop_to_showdown / 2)
    timer.observe(TURN)
    timer.observe(RIVER)
    clock.advance(flop_to_showdown / 2)
    timer.observe(COMPLETE)


def test_a_round_is_timed_from_the_deal_to_the_showdown(timer, clock):
    play_a_round(timer, clock, deal_to_flop=6.0, flop_to_showdown=9.0)

    timings = timer.timings()
    assert timings["deal_to_flop_seconds"] == 6.0
    assert timings["flop_to_showdown_seconds"] == 9.0
    assert timings["round_seconds"] == 15.0


def test_the_start_time_is_recorded_in_wall_clock(timer, clock):
    play_a_round(timer, clock)
    assert timer.timings()["round_started_at"] == "wall-1000.0"


def test_the_gap_to_the_next_round_is_measured_from_the_showdown(timer, clock):
    play_a_round(timer, clock, deal_to_flop=5.0, flop_to_showdown=5.0)
    assert timer.timings()["seconds_since_previous_round"] is None  # first round

    clock.advance(4.0)                     # table clears
    timer.observe(WAITING)
    clock.advance(12.0)                    # waiting for the next deal
    play_a_round(timer, clock, deal_to_flop=3.0, flop_to_showdown=4.0)

    timings = timer.timings()
    assert timings["seconds_since_previous_round"] == 16.0
    assert timings["round_seconds"] == 7.0


def test_each_milestone_is_taken_the_first_time_it_is_seen(timer, clock):
    """Polling repeats the same state; the timing must not drift with it."""
    timer.observe(PLAYER_CARDS)
    for _ in range(5):
        clock.advance(0.35)
        timer.observe(PLAYER_CARDS)

    clock.advance(1.0)
    timer.observe(FLOP)
    flop_at = timer.flop_at
    for _ in range(10):
        clock.advance(0.35)
        timer.observe(FLOP)
    assert timer.flop_at == flop_at

    timer.observe(COMPLETE)
    showdown = timer.showdown_at
    for _ in range(10):
        clock.advance(0.35)
        timer.observe(COMPLETE)
    assert timer.showdown_at == showdown


def test_a_round_that_never_shows_a_flop_reports_what_it_saw(timer, clock):
    timer.observe(PLAYER_CARDS)
    clock.advance(3.0)

    timings = timer.timings()
    assert timings["round_seconds"] is None
    assert timings["deal_to_flop_seconds"] is None
    assert timings["round_started_at"] is not None


def test_the_timer_restarts_when_the_table_empties(timer, clock):
    play_a_round(timer, clock)
    first_start = timer.started_at

    clock.advance(5.0)
    timer.observe(WAITING)
    clock.advance(5.0)
    timer.observe(PLAYER_CARDS)

    assert timer.started_at != first_start
    assert timer.flop_at is None
    assert timer.showdown_at is None


def test_a_new_memory_generation_also_starts_a_new_round(timer, clock):
    """A new deal that never showed an empty table still counts as new."""
    timer.observe(PLAYER_CARDS, generation=1)
    clock.advance(8.0)
    timer.observe(COMPLETE, generation=1)
    assert timer.timings()["round_seconds"] == 8.0

    clock.advance(10.0)
    timer.observe(PLAYER_CARDS, generation=2)
    clock.advance(4.0)
    timer.observe(COMPLETE, generation=2)

    timings = timer.timings()
    assert timings["round_seconds"] == 4.0
    assert timings["seconds_since_previous_round"] == 10.0


def test_timings_are_never_negative(timer, clock):
    """A clock that jumps backwards must not produce a negative round."""
    timer.observe(PLAYER_CARDS)
    clock.now -= 30.0
    timer.observe(COMPLETE)
    assert timer.timings()["round_seconds"] == 0.0


def test_the_log_line_reads_sensibly(timer, clock):
    assert timer.describe() == "no timings yet"
    play_a_round(timer, clock, deal_to_flop=6.0, flop_to_showdown=9.0)
    described = timer.describe()
    assert "round 15.0s" in described
    assert "deal->flop 6.0s" in described


# -- what reaches the spreadsheet ---------------------------------------------

def test_the_timings_reach_the_stored_row():
    from database.db import record_to_row

    record = {
        "player_1": "8S", "player_2": "9C",
        "flop_1": "2D", "flop_2": "QH", "flop_3": "3D",
        "turn": "3H", "river": "6S", "dealer_1": "8D", "dealer_2": "QD",
        "player_hand": "Pair", "dealer_hand": "Two Pair",
        "winner": "Dealer", "dealer_qualified": True,
        "hand_fingerprint": "8S-9C-2D-QH-3D-3H-6S-8D-QD",
        "round_started_at": "2026-09-08 11:00:00",
        "deal_to_flop_seconds": 6.0, "flop_to_showdown_seconds": 9.0,
        "round_seconds": 15.0, "seconds_since_previous_round": 16.0,
    }
    row = record_to_row(record)
    assert row["round_seconds"] == 15.0
    assert row["seconds_since_previous_round"] == 16.0


def test_the_timings_reach_the_spreadsheet(tmp_path):
    import datetime

    from openpyxl import load_workbook

    from export import excel_export

    path = str(tmp_path / "timings.xlsx")
    excel_export.append_hand({
        "id": 1, "recorded_at": datetime.datetime(2026, 9, 8, 11, 0, 15),
        "player_card_1": "8S", "player_card_2": "9C",
        "flop_card_1": "2D", "flop_card_2": "QH", "flop_card_3": "3D",
        "turn_card": "3H", "river_card": "6S",
        "dealer_card_1": "8D", "dealer_card_2": "QD",
        "player_hand": "Pair", "dealer_hand": "Two Pair",
        "winner": "Dealer", "dealer_qualified": True,
        "round_started_at": datetime.datetime(2026, 9, 8, 11, 0, 0),
        "deal_to_flop_seconds": 6.0, "flop_to_showdown_seconds": 9.0,
        "round_seconds": 15.0, "seconds_since_previous_round": 16.0,
    }, path)

    workbook = load_workbook(path)
    headers, row = list(workbook.active.iter_rows(values_only=True))[:2]
    workbook.close()

    columns = dict(zip(headers, row))
    assert columns["Round Started"] == "2026-09-08 11:00:00"
    assert columns["Deal To Flop (s)"] == 6.0
    assert columns["Flop To Showdown (s)"] == 9.0
    assert columns["Round Length (s)"] == 15.0
    assert columns["Since Previous Round (s)"] == 16.0
