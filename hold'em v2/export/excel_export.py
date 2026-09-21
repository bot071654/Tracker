"""Excel mirror of the poker_hands table (data/poker_hands.xlsx).

PostgreSQL stays the source of truth; Excel is a convenience copy that is
appended to after every successful insert, and can be rebuilt from scratch.
"""

import logging
import os

from openpyxl import Workbook, load_workbook

from config.settings import EXCEL_PATH

logger = logging.getLogger(__name__)

HEADERS = [
    "ID", "Recorded At",
    "Player Card 1", "Player Card 2",
    "Flop Card 1", "Flop Card 2", "Flop Card 3",
    "Turn", "River",
    "Dealer Card 1", "Dealer Card 2",
    "Player Hand", "Dealer Hand",
    "Winner", "Dealer Qualified",
    "Round Started", "Deal To Flop (s)", "Flop To Showdown (s)",
    "Round Length (s)", "Since Previous Round (s)",
]

# Excel column -> database column.
_FIELDS = [
    "id", "recorded_at",
    "player_card_1", "player_card_2",
    "flop_card_1", "flop_card_2", "flop_card_3",
    "turn_card", "river_card",
    "dealer_card_1", "dealer_card_2",
    "player_hand", "dealer_hand",
    "winner", "dealer_qualified",
    "round_started_at", "deal_to_flop_seconds", "flop_to_showdown_seconds",
    "round_seconds", "seconds_since_previous_round",
]


class ExcelError(RuntimeError):
    """Raised when the workbook cannot be read or written (often: open in Excel)."""


def _row_values(row):
    values = []
    for field in _FIELDS:
        value = row.get(field)
        if field in ("recorded_at", "round_started_at") and value is not None:
            value = value.strftime("%Y-%m-%d %H:%M:%S") if hasattr(value, "strftime") else str(value)
        elif field == "dealer_qualified" and value is not None:
            value = "Yes" if value else "No"
        values.append(value)
    return values


def _new_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hands"
    sheet.append(HEADERS)
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = \
            max(12, len(header) + 2)
    return workbook


def _save(workbook, path):
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        workbook.save(path)
    except PermissionError as exc:
        raise ExcelError(
            "Cannot write %s - close the file in Excel and try again." % path
        ) from exc
    except OSError as exc:
        raise ExcelError("Cannot write %s - %s" % (path, exc)) from exc


def ensure_workbook(path=EXCEL_PATH):
    """Create the workbook with headers if it does not exist yet."""
    if os.path.exists(path):
        return path
    _save(_new_workbook(), path)
    logger.info("Created Excel workbook %s", path)
    return path


def existing_ids(path=EXCEL_PATH):
    """Set of hand IDs already present in the workbook."""
    if not os.path.exists(path):
        return set()
    try:
        workbook = load_workbook(path, read_only=True)
    except (PermissionError, OSError) as exc:
        raise ExcelError("Cannot read %s - %s" % (path, exc)) from exc
    try:
        sheet = workbook.active
        ids = set()
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
            if row and row[0] is not None:
                ids.add(row[0])
        return ids
    finally:
        workbook.close()


def append_hand(row, path=EXCEL_PATH):
    """Append one stored hand. Returns False if that ID is already present."""
    ensure_workbook(path)
    if row.get("id") in existing_ids(path):
        logger.info("Hand #%s already in Excel, not appended again", row.get("id"))
        return False

    try:
        workbook = load_workbook(path)
    except (PermissionError, OSError) as exc:
        raise ExcelError("Cannot open %s - %s" % (path, exc)) from exc

    workbook.active.append(_row_values(row))
    _save(workbook, path)
    logger.info("Hand #%s appended to Excel", row.get("id"))
    return True


def rebuild(rows, path=EXCEL_PATH):
    """Rewrite the workbook from the given database rows. Returns the count."""
    workbook = _new_workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(_row_values(row))
    _add_statistics(workbook, rows)
    _save(workbook, path)
    logger.info("Rebuilt Excel workbook with %d hand(s): %s", len(rows), path)
    return len(rows)


def _add_statistics(workbook, rows):
    """A second sheet summarising how the recorded rounds turned out.

    Written only on a rebuild, where every row is to hand. Appending a single
    hand leaves it alone rather than rewriting a summary on every insert - the
    Statistics sheet is as current as the last Export Excel.
    """
    from database.db import running_win_percentages, summarise_results

    counts = {}
    for row in rows:
        winner = row.get("winner")
        if winner:
            counts[winner] = counts.get(winner, 0) + 1
    stats = summarise_results(counts)

    sheet = workbook.create_sheet("Statistics")
    sheet.append(["Results across every recorded round"])
    sheet.append([])
    for label, value in (
        ("Total rounds", stats["rounds"]),
        ("Player wins", stats["player"]),
        ("Dealer wins", stats["dealer"]),
        ("Ties", stats["tie"]),
    ):
        sheet.append([label, value])
    sheet.append([])
    for label, value in (
        ("Player win %", stats["player_percent"]),
        ("Dealer win %", stats["dealer_percent"]),
        ("Tie %", stats["tie_percent"]),
    ):
        sheet.append([label, value])

    sheet.append([])
    sheet.append(["The player's win rate as it developed"])
    sheet.append(["Round", "Winner", "Player wins so far", "Player win %"])
    for entry in running_win_percentages(rows):
        sheet.append([entry["round"], entry["winner"],
                      entry["player_wins"], entry["player_percent"]])

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 14
